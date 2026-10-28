"""Build MAIN_MANUSCRIPT_JQL_SUBMIT_FINAL.docx from MANUSCRIPT.md.

- {key} citation tokens -> numbered [n] by first appearance (Vancouver order)
- figures/tables inserted inline right after the paragraph of first citation
- OMML math for $...$ segments; TNR 12pt body, 10pt captions
- numbers are taken from generated tables/*.csv, not retyped
"""
import csv
import json
import os
import re
import sys

sys.path.insert(0, os.path.expanduser(
    "~/repos/wip/democratic-dictator"))
from docx_math import add_text_to_para  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt, Inches  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

ROOT = os.path.join(os.path.dirname(__file__), "..")
MS = os.path.join(ROOT, "manuscript", "MANUSCRIPT.md")
OUT = os.path.join(ROOT, "manuscript", "MAIN_MANUSCRIPT_JQL_SUBMIT_FINAL.docx")

REFS = {
    "daniels": "Daniels PT. An Exploration of Writing. Sheffield: Equinox; 2018.",
    "sproat": "Sproat R. A Computational Theory of Writing Systems. Cambridge: Cambridge University Press; 2000.",
    "changizi": "Changizi MA, Shimojo S. Character complexity and redundancy in writing systems over human history. Proc R Soc B. 2005;272(1560):267-275.",
    "miton": "Miton H, Morin O. Graphic complexity in writing systems. Cognition. 2021;214:104771. doi:10.1016/j.cognition.2021.104771",
    "coupe": "Coupe C, Oh YM, Dediu D, Pellegrino F. Different languages, similar encoding efficiency: comparable information rates across the human communicative niche. Sci Adv. 2019;5(9):eaaw2594. doi:10.1126/sciadv.aaw2594",
    "jaeger": "Jaeger TF, Graff P, Croft W, Pontillo D. Mixed effect models for genetic and areal dependencies in linguistic typology. Linguist Typol. 2011;15(2):281-319.",
    "bickel": "Bickel B. Typology in the 21st century: major current developments. Linguist Typol. 2007;11(1):239-251.",
    "phoible": "Moran S, McCloy D, editors. PHOIBLE 2.0. Jena: Max Planck Institute for the Science of Human History; 2019. https://phoible.org",
    "moran_phoneme": "Moran S, Grossman E, Verkerk A. Investigating diachronic trends in phonological inventories using BDPROTO. Lang Resour Eval. 2021;55(1):79-103. doi:10.1007/s10579-019-09483-3",
    "cldr": "Unicode Consortium. Unicode Common Locale Data Repository (CLDR), version 46.1; 2024. https://cldr.unicode.org",
    "glottolog": "Hammarstrom H, Forkel R, Haspelmath M, Bank S. Glottolog 5.0. Leipzig: Max Planck Institute for Evolutionary Anthropology; 2024. https://glottolog.org",
    "asjp": "Wichmann S, Holman EW, Brown CH. The ASJP Database (version 20). 2022. https://asjp.clld.org",
    "wordfreq": "Speer R. rspeer/wordfreq: v3.0 [software]. Zenodo; 2022. https://doi.org/10.5281/zenodo.7199437",
    "maddieson": "Maddieson I. Patterns of Sounds. Cambridge: Cambridge University Press; 1984.",
    "fenk": "Fenk-Oczlon G, Fenk A. Complexity trade-offs between the subsystems of language. In: Miestamo M, Sinnemaki K, Karlsson F, editors. Language Complexity. Amsterdam: Benjamins; 2008. p. 43-65.",
}

# literal labels as cited in MANUSCRIPT.md (already first-citation order)
TABDEFS = {
    "Table 1": ("table_primary_class_medians.csv",
                "Class medians of G, P, and G/P with script-cluster bootstrap 95% CIs."),
    "Table 2": ("table_primary_within_class_rho.csv",
                "Within-class Spearman rho(P, G) with family-cluster bootstrap CIs."),
    "Table 3": ("table_primary_model.csv",
                "Crossed mixed-effects model logG ~ logP x C(granularity), fixed effects."),
    "Table 4": ("table_simulation_bounds.csv",
                "Simulated stored-unit vs compositional bounds vs observed G."),
    "Table 5": ("table_simulation_counterfactual.csv",
                "Counterfactual stored-unit demand by granularity."),
    "Table 6": ("table_escape_spearman.csv",
                "Combinatorial escape: Spearman rho of inventory vs mean word length."),
    "Table 7": ("table_inverse_combined.csv",
                "Inverse inference: macro-F1 under grouped cross-validation."),
}
FIGDEFS = {
    "Figure 1": ("fig_scatter_and_medians.png",
                 "P vs G across writing-system classes with class medians."),
    "Figure 2": ("fig_simulation.png",
                 "Observed G vs simulated stored-unit demand by granularity."),
    "Figure 3": ("fig_escape.png",
                 "Spoken word length vs P and G (ASJP)."),
}


def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_para(doc, text, size=12, bold=False, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    add_text_to_para(p, text, bold=bold, font_size=size)
    for r in p.runs:
        r.font.name = "Times New Roman"
    return p


HEADER_MAP = {
    "gran": "granularity", "var": "variable", "n_scripts": "scripts (n)",
    "median": "median", "ci_lo": "CI lo", "ci_hi": "CI hi", "n": "n",
    "ws_class": "class", "n_families": "families (n)", "rho": "rho",
    "p": "p", "term": "term", "coef": "coef",
    "granularity": "granularity", "obs_G_median": "obs. G median",
    "sim_stored_median": "stored bound (median)",
    "sim_compositional_median": "compositional bound (median)",
    "obs_within_bounds_frac": "within bounds (frac)",
    "reencoded_granularity": "re-encoded granularity",
    "U_median": "U median", "U_iqr": "U IQR", "range": "range",
    "subset": "subset", "rho_P_vs_wlen": "rho(P, word len.)",
    "rho_P_vs_wlen_p": "p", "rho_P_vs_wlen_ci": "95% CI",
    "rho_G_vs_wlen": "rho(G, word len.)", "rho_G_vs_wlen_p": "p",
    "rho_G_vs_wlen_ci": "95% CI", "model": "model", "target": "target",
    "mean_macroF1": "macro-F1", "sd": "SD", "control": "control",
    "observed": "observed", "null_95": "null 95%",
    "perm_mean": "perm. mean", "perm_sd": "perm. SD", "emp_p": "emp. p",
    "mean_word_segments": "word len.",
}


def fmt_cell(val, header):
    v = val.strip()
    # confidence intervals like [-0.086,0.282] -> [-0.09, 0.28]
    m = re.fullmatch(r"\[(-?\d+\.?\d*),(-?\d+\.?\d*)\]", v)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        if lo == hi:  # degenerate bootstrap CI (single-family class)
            return "not estimable"
        return f"[{lo:.2f}, {hi:.2f}]"
    try:
        f = float(v)
    except ValueError:
        return {"GP_ratio": "G/P", "P_total": "P"}.get(v, v)
    h = header.lower()
    if f != f:  # NaN
        return "not estimable" if "ci" in h else "-"
    if "_p" in h or h == "p" or h == "emp. p":
        return "<0.001" if f < 0.001 else f"{f:.3f}"
    if f == int(f) and abs(f) < 1e7:
        return str(int(f))
    if "f1" in h or "rho" in h or "coef" in h or "perm" in h or h == "sd" \
            or "ci" in h or "frac" in h or "bound" in h or "median" in h \
            or "len" in h or "observed" in h or "null" in h \
            or h in ("median", "u iqr", "range"):
        return f"{f:.2f}" if abs(f) < 10 else f"{f:.1f}" if abs(f) < 1000 \
            else f"{f:.0f}"
    return f"{f:.3f}" if abs(f) < 100 else f"{f:.0f}"


TABLE_NOTES = {
    "table_primary_within_class_rho.csv":
        "family-cluster confidence intervals are not estimable for classes "
        "represented by a single language family.",
    "table_escape_spearman.csv":
        "p values are conventional Spearman test p values; 95% CIs are "
        "family-cluster bootstrap intervals. Because the confidence intervals "
        "account for genealogical clustering while the nominal p values do not, "
        "the two inferential summaries need not coincide. Cluster-aware "
        "confidence intervals are preferred for interpretation. Family-cluster "
        "confidence intervals are not estimable for classes represented by a "
        "single language family.",
}


def add_table(doc, csv_path, caption, tnum):
    rows = list(csv.reader(open(csv_path)))
    headers = [HEADER_MAP.get(h, h.replace("_", " ")) for h in rows[0]]
    cp = doc.add_paragraph()
    r = cp.add_run(f"Table {tnum}. {caption}")
    r.font.size = Pt(10); r.font.name = "Times New Roman"; r.italic = True
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            text = headers[j] if i == 0 else fmt_cell(cell, headers[j])
            run = c.paragraphs[0].add_run(text)
            run.font.size = Pt(10)
            run.font.name = "Times New Roman"
            if i == 0:
                run.bold = True
    note = TABLE_NOTES.get(os.path.basename(csv_path))
    if note:
        np_ = doc.add_paragraph()
        nr = np_.add_run(f"Note: {note}")
        nr.font.size = Pt(10); nr.font.name = "Times New Roman"
        nr.italic = True
    doc.add_paragraph()


def add_figure(doc, fig_path, caption, fnum):
    doc.add_picture(fig_path, width=Inches(5.8))
    cp = doc.add_paragraph()
    r = cp.add_run(f"Figure {fnum}. {caption}")
    r.font.size = Pt(10); r.font.name = "Times New Roman"; r.italic = True


REGISTRY = json.load(open(os.path.join(ROOT, "results", "canonical_results.json")))


def subst_tokens(text):
    def rep(m):
        k = m.group(1)
        assert k in REGISTRY, f"unknown token {{{k}}}"
        return REGISTRY[k]
    text = re.sub(r"\{([A-Z][A-Z0-9_]+)\}", rep, text)
    assert not re.search(r"\{[A-Z][A-Z0-9_]+\}", text), "unresolved token remains"
    return text


def main():
    text = open(MS).read().replace("{refs}", "@@REFS@@")
    text = subst_tokens(text)

    # assign citation numbers by first appearance
    order, cite_map = [], {}
    def repl(m):
        for k in m.group(1).split(", "):
            k = k.strip()
            if k not in cite_map:
                order.append(k); cite_map[k] = len(order)
        return "[" + ",".join(str(cite_map[k.strip()]) for k in m.group(1).split(", ")) + "]"
    text = re.sub(r"\{([a-z_]+(?:, [a-z_]+)*)\}", repl, text)
    missing = [k for k in order if k not in REFS]
    assert not missing, f"missing refs: {missing}"

    # verify every declared object is actually cited, in first-citation order
    cited_order = re.findall(r"\b(?:Table|Figure) \d\b", text)
    for label in TABDEFS | FIGDEFS:
        assert label in cited_order, f"{label} declared but not cited"
    nums = {"Table": [], "Figure": []}
    for lab in cited_order:
        kind, n = lab.split()
        nums[kind].append(int(n))
    for kind, seq in nums.items():
        first_seen = list(dict.fromkeys(seq))
        assert first_seen == list(range(1, len(first_seen) + 1)), (
            f"{kind} citations out of first-citation order or gapped: {seq}")

    doc = Document()
    style_doc(doc)
    lines = text.split("\n")
    inserted = set()
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if ln.startswith("# "):
            add_para(doc, ln[2:], size=14, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER)
        elif ln.startswith("## "):
            add_para(doc, ln[3:], size=12, bold=True)
        elif ln.startswith("### "):
            add_para(doc, ln[4:], size=12, bold=True)
        elif ln == "@@REFS@@":
            for i, k in enumerate(order, 1):
                p = doc.add_paragraph()
                r = p.add_run(f"{i}. {REFS[k]}")
                r.font.size = Pt(11); r.font.name = "Times New Roman"
        else:
            ln = re.sub(r"\*\*([^*]+)\*\*", r"\1", ln)
            add_para(doc, ln)
        # inline objects right after first-citation paragraph
        for label, (fn, cap) in TABDEFS.items():
            if label in ln and label not in inserted:
                add_table(doc, os.path.join(ROOT, "tables", fn), cap,
                          int(label.split()[-1]))
                inserted.add(label)
        for label, (fn, cap) in FIGDEFS.items():
            if label in ln and label not in inserted:
                add_figure(doc, os.path.join(ROOT, "figures", fn), cap,
                           int(label.split()[-1]))
                inserted.add(label)
    uninserted = (set(TABDEFS) | set(FIGDEFS)) - inserted
    assert not uninserted, f"uncited objects: {uninserted}"
    doc.save(OUT)
    print("wrote", OUT, "| refs:", len(order), "| inserted:", sorted(inserted))


if __name__ == "__main__":
    main()
