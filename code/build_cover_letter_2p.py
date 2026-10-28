import json
import os
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.join(os.path.dirname(__file__), "..")
REGISTRY = json.load(open(os.path.join(ROOT, "results", "canonical_results.json")))

paras = []
with open(os.path.join(ROOT, 'manuscript/COVER_LETTER.md'), encoding='utf-8') as f:
    text = f.read().strip()

# resolve {TOKEN} placeholders from the canonical registry; the cover letter
# uses ASCII minus signs, so convert the typographic minus in values.
def rep(m):
    return REGISTRY[m.group(1)].replace("−", "-")
text = re.sub(r"\{([A-Z][A-Z0-9_]+)\}", rep, text)
assert not re.search(r"\{[A-Z][A-Z0-9_]+\}", text), "unresolved token remains"
blocks = [b.strip() for b in text.split('\n\n') if b.strip()]

d = Document()
st = d.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)

sec = d.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.4)
sec.left_margin = sec.right_margin = Cm(2.4)

BREAK_BEFORE = "The paper's second contribution"

for i, b in enumerate(blocks):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(14)
    pf.widow_control = True
    lines = b.split('\n')
    for j, ln in enumerate(lines):
        if j:
            p.add_run().add_break()
        p.add_run(ln)
    if b.startswith(BREAK_BEFORE):
        pf.page_break_before = True
    if i >= len(blocks) - 2:  # Sincerely + signature block kept together
        pf.keep_with_next = True
        pf.keep_together = True
        pf.space_before = Pt(10) if b.startswith('Sincerely') else Pt(0)

d.save(os.path.join(ROOT, 'manuscript/COVER_LETTER_JQL_SUBMIT_FINAL_2P.docx'))
print('saved', len(blocks), 'blocks')
